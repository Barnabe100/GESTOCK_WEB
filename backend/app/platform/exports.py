"""Exports de données (Lot 2, ADR-0038) : UNE action « Exporter », format choisi par
l'utilisateur parmi ceux que la fonctionnalité propose.

Principe : la fonctionnalité sélectionne ses données UNE seule fois (mêmes filtres, même
périmètre tenant / site / permission / portée que sa liste) et produit un ``ExportTable`` ;
les formats ne changent que la représentation, jamais le périmètre :

    filtres → portée + permissions → requête commune → ExportTable → CSV | Excel | PDF

- CSV : séparateur « ; », UTF-8 avec BOM, décimales à virgule (Excel en français).
- Excel (.xlsx) : nombres et dates typés, en-tête figé, filtre automatique.
- PDF : A4 paysage, adapté à l'impression (rapport) ; jamais le format des reçus (80 mm).

Tout export est audité (``export.generated``) : fonctionnalité, format, filtres réellement
utilisés (aucune valeur inventée), nombre de lignes.
"""

import csv
import io
import uuid
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from enum import StrEnum
from typing import Any
from zoneinfo import ZoneInfo

from fastapi import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.errors import BusinessRuleError
from app.platform.audit.service import audit_action
from app.platform.context import RequestContext


class ExportFormat(StrEnum):
    XLSX = "xlsx"
    CSV = "csv"
    PDF = "pdf"


MEDIA_TYPES = {
    ExportFormat.XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ExportFormat.CSV: "text/csv; charset=utf-8",
    ExportFormat.PDF: "application/pdf",
}


class ColumnKind(StrEnum):
    TEXT = "text"
    MONEY = "money"
    QUANTITY = "quantity"
    DATE = "date"
    DATETIME = "datetime"


@dataclass(frozen=True)
class ExportColumn:
    header: str
    kind: ColumnKind = ColumnKind.TEXT
    # Largeur relative (PDF) et indicative (Excel).
    width: int = 12


@dataclass
class ExportTable:
    """Jeu de données autorisé, déjà sélectionné par la fonctionnalité (valeurs brutes :
    ``Decimal`` pour les montants, ``date`` / ``datetime`` aware pour les dates)."""

    title: str
    file_stem: str
    columns: Sequence[ExportColumn]
    rows: list[list[Any]]
    timezone: str
    # Lignes d'en-tête du PDF (entreprise, date de génération, filtres appliqués).
    context_lines: list[str] = field(default_factory=list)


def check_format(fmt: ExportFormat, allowed: Sequence[ExportFormat]) -> None:
    """Formats proposés par la fonctionnalité : aucun autre n'est accepté."""
    if fmt not in allowed:
        raise BusinessRuleError(
            "Format d'export non disponible pour cette fonctionnalité",
            code="export_format_unavailable",
            extra={"allowed": [f.value for f in allowed]},
        )


def check_row_limit(count: int, limit: int) -> None:
    if count > limit:
        raise BusinessRuleError(
            "Trop de lignes à exporter : affinez les filtres",
            code="export_too_large",
            extra={"count": count, "limit": limit},
        )


# --- Mise en forme des valeurs ------------------------------------------------------------


def _local(value: datetime, tz: str) -> datetime:
    return value.astimezone(ZoneInfo(tz)) if value.tzinfo else value


def _fr_number(value: Decimal, places: int) -> str:
    """« 1234,50 » (CSV) : nombre reconnu par Excel en français, sans séparateur de milliers."""
    return f"{value:.{places}f}".replace(".", ",")


def _fr_grouped(value: Decimal, places: int) -> str:
    """« 1 234,50 » (PDF), espace simple (police standard)."""
    return f"{value:,.{places}f}".replace(",", " ").replace(".", ",")


def _text(value: Any, kind: ColumnKind, tz: str, *, grouped: bool) -> str:
    if value is None:
        return ""
    if kind is ColumnKind.MONEY and isinstance(value, Decimal):
        return _fr_grouped(value, 2) if grouped else _fr_number(value, 2)
    if kind is ColumnKind.QUANTITY and isinstance(value, Decimal):
        return _fr_grouped(value, 3) if grouped else _fr_number(value, 3)
    if kind is ColumnKind.DATETIME and isinstance(value, datetime):
        return _local(value, tz).strftime("%d/%m/%Y %H:%M")
    if kind is ColumnKind.DATE and isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return str(value)


# --- Rendus -------------------------------------------------------------------------------


def render_csv(table: ExportTable) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow([c.header for c in table.columns])
    for row in table.rows:
        writer.writerow(
            [
                _text(v, c.kind, table.timezone, grouped=False)
                for v, c in zip(row, table.columns, strict=True)
            ]
        )
    return buffer.getvalue().encode("utf-8-sig")


_EXCEL_FORMATS = {
    ColumnKind.MONEY: "#,##0.00",
    ColumnKind.QUANTITY: "#,##0.000",
    ColumnKind.DATE: "dd/mm/yyyy",
    ColumnKind.DATETIME: "dd/mm/yyyy hh:mm",
}


def _excel_value(value: Any, kind: ColumnKind, tz: str) -> Any:
    if isinstance(value, datetime):
        # Excel n'a pas de fuseau : heure locale de l'entreprise (ADR-0028).
        return _local(value, tz).replace(tzinfo=None)
    if isinstance(value, uuid.UUID):
        return str(value)
    return value


def render_xlsx(table: ExportTable) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = table.title[:31]
    sheet.append([c.header for c in table.columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in table.rows:
        sheet.append(
            [
                _excel_value(v, c.kind, table.timezone)
                for v, c in zip(row, table.columns, strict=True)
            ]
        )
    for index, column in enumerate(table.columns, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = max(column.width, len(column.header) + 2)
        number_format = _EXCEL_FORMATS.get(column.kind)
        if number_format:
            for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = number_format
    sheet.freeze_panes = "A2"
    if table.rows:
        sheet.auto_filter.ref = sheet.dimensions
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_pdf(table: ExportTable) -> bytes:
    """Rapport A4 paysage : titre, contexte, tableau (en-tête répété sur chaque page),
    pagination."""
    buffer = io.BytesIO()
    page = landscape(A4)
    margin = 12 * mm
    document = SimpleDocTemplate(
        buffer,
        pagesize=page,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
        title=table.title,
    )
    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"].clone("cell", fontSize=7.5, leading=9)
    head_style = cell_style.clone("head", fontName="Helvetica-Bold")
    story: list[Any] = [Paragraph(table.title, styles["Title"])]
    for line in table.context_lines:
        story.append(Paragraph(_escape(line), styles["Normal"]))
    story.append(Spacer(1, 4 * mm))
    total_width = page[0] - 2 * margin
    weights = sum(c.width for c in table.columns)
    widths = [total_width * c.width / weights for c in table.columns]
    data: list[list[Any]] = [[Paragraph(_escape(c.header), head_style) for c in table.columns]]
    for row in table.rows:
        data.append(
            [
                Paragraph(_escape(_text(v, c.kind, table.timezone, grouped=True)), cell_style)
                for v, c in zip(row, table.columns, strict=True)
            ]
        )
    grid = Table(data, colWidths=widths, repeatRows=1)
    style: list[Any] = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B8C2CC")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]
    for index, column in enumerate(table.columns):
        if column.kind in (ColumnKind.MONEY, ColumnKind.QUANTITY):
            style.append(("ALIGN", (index, 1), (index, -1), "RIGHT"))
    grid.setStyle(TableStyle(style))
    story.append(grid)

    def footer(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawRightString(page[0] - margin, margin / 2, f"Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


RENDERERS = {
    ExportFormat.CSV: render_csv,
    ExportFormat.XLSX: render_xlsx,
    ExportFormat.PDF: render_pdf,
}


def export_response(table: ExportTable, fmt: ExportFormat, now: datetime) -> Response:
    content = RENDERERS[fmt](table)
    stamp = _local(now, table.timezone).strftime("%Y%m%d-%H%M")
    filename = f"{table.file_stem}-{stamp}.{fmt.value}"
    return Response(
        content=content,
        media_type=MEDIA_TYPES[fmt],
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-store",
        },
    )


# --- Audit -----------------------------------------------------------------------------------


def audit_export(
    db: Any,
    ctx: RequestContext,
    *,
    feature: str,
    fmt: ExportFormat,
    filters: dict[str, Any],
    row_count: int,
    site_id: uuid.UUID | None = None,
) -> None:
    """Trace d'un export (même transaction que la lecture) : utilisateur, tenant et date sont
    portés par l'entrée d'audit ; ``filters`` ne contient que les filtres renseignés."""
    audit_action(
        db,
        ctx,
        "export.generated",
        entity_type="export",
        entity_id=None,
        site_id=site_id,
        data={
            "feature": feature,
            "format": fmt.value,
            "filters": filters,
            "row_count": row_count,
        },
    )
