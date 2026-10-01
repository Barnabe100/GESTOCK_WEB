"""Lot 2 — historique des ventes : tri chronologique par défaut, filtres (vendeur / opérateur,
« Mes ventes », article, référence article, référence de paiement, client, canal), exports
Excel / CSV / PDF limités au périmètre de la liste et audités, permission
``sales.sale.export``, limite de crédit (``customers.credit_limit.manage``), mouvements de stock
de la vente (``stock.movement.view``) et chronologie (``audit.log.view``)."""

import csv
import io
import uuid
from datetime import UTC, datetime
from decimal import Decimal
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.config import Settings
from tests import stock_helpers as sh
from tests.conftest import PASSWORD, Api, login
from tests.stock_helpers import World

YEAR = datetime.now(UTC).year


@pytest.fixture
def priced(world: World, owner_db: Session) -> World:
    """Article 0 à 10 000 et article 1 à 2 500 ; stock sur les deux sites."""
    for index, price in ((0, 10000), (1, 2500)):
        owner_db.execute(
            text("UPDATE catalog_articles SET sale_price = :p WHERE id = :id"),
            {"p": price, "id": world.articles[index]},
        )
    owner_db.commit()
    sh.validated_entry(world, [(0, "200", "6000"), (1, "100", "1000")])
    sh.validated_entry(world, [(0, "100", "6000")], site_id=world.site2)
    return world


# --- Aides ----------------------------------------------------------------------------------------


def _sale(
    w: World,
    api: Api | None = None,
    *,
    article: int = 0,
    units: int = 1,
    site: str | None = None,
    payments: list[dict[str, Any]] | None = None,
    validate: bool = True,
    **extra: Any,
) -> dict[str, Any]:
    api = api or w.owner
    created = api.post(
        "/sales",
        json={
            "site_id": site or w.site,
            "lines": [{"article_id": w.articles[article], "quantity": str(units)}],
            **extra,
        },
    )
    assert created.status_code == 201, created.text
    sale = dict(created.json())
    if not validate:
        return sale
    if payments is None:
        payments = [{"amount": sale["total"], "method": "CASH"}]
    validated = api.post(f"/sales/{sale['id']}/validate", json={"payments": payments})
    assert validated.status_code == 200, validated.text
    return dict(validated.json())


def _custom_member(
    w: World, client: TestClient, email: str, permissions: list[str], **access: Any
) -> Api:
    role = w.owner.post("/roles", json={"name": f"Rôle {email}", "permissions": permissions})
    assert role.status_code == 201, role.text
    created = w.owner.post(
        "/members",
        json={
            "email": email,
            "full_name": email,
            "password": "Provisoire-123",
            "roles": [{"role_id": role.json()["id"]}],
            **(access or {"all_sites": True}),
        },
    )
    assert created.status_code == 201, created.text
    token = login(client, email, "Provisoire-123").json()["access_token"]
    Api(client, token).post(
        "/me/password", json={"current_password": "Provisoire-123", "new_password": PASSWORD}
    )
    return Api(client, login(client, email).json()["access_token"])


def _ids(api: Api, **params: Any) -> list[str]:
    response = api.get("/sales", params={"limit": 100, **params})
    assert response.status_code == 200, response.text
    return [str(s["id"]) for s in response.json()["items"]]


def _export(api: Api, fmt: str, **params: Any) -> Any:
    return api.get("/sales/export", params={"format": fmt, **params})


def _csv_rows(response: Any) -> list[list[str]]:
    assert response.content.startswith(b"\xef\xbb\xbf")  # BOM UTF-8 (Excel en français)
    return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"))


def _export_logs(w: World) -> list[dict[str, Any]]:
    items = w.owner.get("/audit-logs", params={"action": "export.generated", "limit": 50})
    return [dict(i) for i in items.json()["items"]]


SELLER_PERMISSIONS = [
    "sales.sale.view",
    "sales.sale.create",
    "sales.sale.validate",
    "sales.payment.create",
]


# --- Tri par défaut -----------------------------------------------------------------------------


def test_default_sort_is_chronological_and_number_sort_stays_available(
    priced: World, owner_db: Session
) -> None:
    first = _sale(priced)
    owner_db.execute(
        text("UPDATE document_sequences SET next_value = 999998 WHERE sequence_key = :k"),
        {"k": f"{priced.site}:sale:{YEAR}"},
    )
    owner_db.commit()
    second = _sale(priced)
    third = _sale(priced)
    code = first["number"].split("-")[1]
    assert second["number"] == f"VENT-{code}-{YEAR}-999999"
    assert third["number"] == f"VENT-{code}-{YEAR}-1000000"
    # Par défaut : la plus récente d'abord, quelle que soit la forme du numéro.
    assert _ids(priced.owner) == [third["id"], second["id"], first["id"]]
    assert _ids(priced.owner, sort="-created_at") == [third["id"], second["id"], first["id"]]
    # Le tri par numéro reste proposé (ordre alphabétique du numéro, option explicite).
    by_number = priced.owner.get("/sales", params={"sort": "number"})
    assert by_number.status_code == 200
    assert len(by_number.json()["items"]) == 3


# --- Filtres --------------------------------------------------------------------------------------


def test_seller_and_mine_filters(priced: World, client: TestClient) -> None:
    seller = sh.member(priced, client, "vendeur@example.com", "seller", all_sites=True)
    manager = sh.member(priced, client, "gestion@example.com", "manager", all_sites=True)
    by_owner = _sale(priced)
    by_seller = _sale(priced, seller)
    by_manager = _sale(priced, manager)
    sellers = {s["name"]: s["id"] for s in priced.owner.get("/sales/sellers").json()}
    assert set(sellers) == {"Owner alpha", "vendeur@example.com", "gestion@example.com"}
    seller_id = sellers["vendeur@example.com"]
    assert _ids(priced.owner, seller_id=seller_id) == [by_seller["id"]]
    # « Mes ventes » : ventes enregistrées par l'utilisateur courant (aucun nom de rôle).
    assert _ids(priced.owner, mine=True) == [by_owner["id"]]
    assert _ids(manager, mine=True) == [by_manager["id"]]
    # Portée « ses propres ventes » : le vendeur ne voit que lui-même, filtre ou pas.
    assert [s["id"] for s in seller.get("/sales/sellers").json()] == [seller_id]
    assert _ids(seller, seller_id=sellers["Owner alpha"]) == []
    assert by_seller["created_by_name"] == "vendeur@example.com"


def test_article_and_reference_filters(priced: World) -> None:
    mobile = priced.owner.post(
        "/payment-methods",
        json={"label": "Orange Money", "kind": "MOBILE_MONEY", "reference_required": True},
    ).json()
    widget = _sale(priced, article=0)
    bolt = _sale(
        priced,
        article=1,
        payments=[{"amount": "2500", "payment_method_id": mobile["id"], "reference": "OM-77812"}],
    )
    assert _ids(priced.owner, article_id=priced.articles[1]) == [bolt["id"]]
    # Référence article (référence / code-barres) ≠ référence de paiement (transaction).
    assert _ids(priced.owner, article_reference="A-0") == [widget["id"]]
    assert _ids(priced.owner, payment_reference="om-778") == [bolt["id"]]
    assert _ids(priced.owner, article_reference="OM-77812") == []
    assert _ids(priced.owner, payment_reference="A-1") == []


def test_customer_and_channel_filters(priced: World) -> None:
    customer = sh.credit_customer(priced)
    with_customer = _sale(priced, customer_id=customer)
    anonymous = _sale(priced)
    pos = priced.owner.post(
        "/pos/checkout",
        json={
            "site_id": priced.site,
            "lines": [{"article_id": priced.articles[0], "quantity": "1"}],
            "payments": [{"amount": "10000", "method": "CASH"}],
            "idempotency_key": str(uuid.uuid4()),
        },
    )
    assert pos.status_code == 201, pos.text
    pos_id = pos.json()["sale"]["id"] if "sale" in pos.json() else pos.json()["id"]
    assert _ids(priced.owner, customer_id=customer) == [with_customer["id"]]
    assert _ids(priced.owner, channel="POS") == [pos_id]
    assert set(_ids(priced.owner, channel="BACKOFFICE")) == {with_customer["id"], anonymous["id"]}


# --- Exports --------------------------------------------------------------------------------------


def test_export_permission(priced: World, client: TestClient) -> None:
    _sale(priced)
    seller = sh.member(priced, client, "vendeur@example.com", "seller", all_sites=True)
    viewer = sh.member(priced, client, "consultant@example.com", "viewer", all_sites=True)
    manager = sh.member(priced, client, "gestion@example.com", "manager", all_sites=True)
    for api in (seller, viewer):
        denied = _export(api, "csv")
        assert denied.status_code == 403 and denied.json()["code"] == "permission_denied"
    assert _export(manager, "csv").status_code == 200
    assert _export(priced.owner, "xlsx").status_code == 200
    # Exporter sans voir les ventes : refusé (l'export ne remplace aucun contrôle).
    blind = _custom_member(priced, client, "aveugle@example.com", ["sales.sale.export"])
    assert _export(blind, "csv").status_code == 403
    unknown = _export(priced.owner, "json")
    assert unknown.status_code == 422


def test_export_formats(priced: World) -> None:
    customer = sh.credit_customer(priced)
    sale = _sale(priced, units=2, customer_id=customer)
    _sale(priced, validate=False)  # brouillon : non numéroté, exporté comme dans la liste

    csv_response = _export(priced.owner, "csv")
    assert csv_response.status_code == 200
    assert csv_response.headers["content-type"].startswith("text/csv")
    disposition = csv_response.headers["content-disposition"]
    assert disposition.startswith('attachment; filename="ventes-') and disposition.endswith('.csv"')
    rows = _csv_rows(csv_response)
    assert rows[0][:3] == ["Numéro", "Date de vente", "Site"]
    assert len(rows) == 3
    exported = next(r for r in rows if r[0] == sale["number"])
    assert exported[3] == "Client comptoir"
    assert exported[7] == "20000,00"  # décimale à virgule, sans séparateur de milliers
    assert exported[6] == "Validée"
    assert any(r[0] == "Non numérotée" and r[6] == "Brouillon" for r in rows[1:])

    xlsx = _export(priced.owner, "xlsx")
    assert xlsx.status_code == 200
    assert xlsx.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    assert sheet is not None
    values = list(sheet.iter_rows(values_only=True))
    assert values[0][0] == "Numéro" and len(values) == 3
    row = next(r for r in values if r[0] == sale["number"])
    assert Decimal(str(row[7])) == Decimal("20000")  # nombre typé, pas du texte
    assert sheet.freeze_panes == "A2"

    pdf = _export(priced.owner, "pdf")
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF")


def test_export_matches_list_scope_and_filters(priced: World, client: TestClient) -> None:
    """Même périmètre que la liste : portée « ses ventes », filtres, sites ; seul le format
    change."""
    own_exporter = _custom_member(
        priced,
        client,
        "export@example.com",
        [*SELLER_PERMISSIONS, "sales.sale.export"],
    )
    mine = _sale(priced, own_exporter)
    _sale(priced)  # vente du propriétaire : hors de la portée « ses propres ventes »
    _sale(priced, own_exporter, site=priced.site2)
    listed = own_exporter.get("/sales").json()
    assert listed["total"] == 2
    for fmt in ("csv", "xlsx"):
        response = _export(own_exporter, fmt)
        assert response.status_code == 200
    rows = _csv_rows(_export(own_exporter, "csv"))[1:]
    assert sorted(r[0] for r in rows) == sorted(s["number"] for s in listed["items"])
    filtered = _csv_rows(_export(own_exporter, "csv", site_id=priced.site))[1:]
    assert [r[0] for r in filtered] == [mine["number"]]
    # Même ordre que la liste (tri par défaut chronologique).
    owner_rows = _csv_rows(_export(priced.owner, "csv"))[1:]
    owner_list = priced.owner.get("/sales").json()["items"]
    assert [r[0] for r in owner_rows] == [s["number"] for s in owner_list]


def test_export_site_and_tenant_isolation(
    priced: World, client: TestClient, provision: Any, api_for: Any
) -> None:
    on_main = _sale(priced)
    on_depot = _sale(priced, site=priced.site2)
    depot_manager = sh.member(
        priced, client, "depot@example.com", "manager", all_sites=False, site_ids=[priced.site2]
    )
    rows = _csv_rows(_export(depot_manager, "csv"))[1:]
    assert [r[0] for r in rows] == [on_depot["number"]]
    # Filtre sur un site non accessible : rien (jamais les ventes de ce site).
    assert _csv_rows(_export(depot_manager, "csv", site_id=priced.site))[1:] == []
    # Autre entreprise : aucune vente d'alpha, même avec le filtre d'un site d'alpha.
    provision("beta")
    beta = api_for("owner@beta.example.com")
    assert _csv_rows(_export(beta, "csv"))[1:] == []
    assert _csv_rows(_export(beta, "csv", site_id=priced.site))[1:] == []
    assert on_main["number"] not in _export(beta, "csv").text


def test_export_requires_permission_on_each_site(priced: World, client: TestClient) -> None:
    """``sales.sale.export`` accordé sur un seul site : les ventes des autres sites visibles ne
    sont jamais exportées."""
    viewer_role = priced.owner.post(
        "/roles", json={"name": "Vue", "permissions": ["sales.sale.view", "sales.sale.view_all"]}
    ).json()
    export_role = priced.owner.post(
        "/roles", json={"name": "Export", "permissions": ["sales.sale.export"]}
    ).json()
    created = priced.owner.post(
        "/members",
        json={
            "email": "multi@example.com",
            "full_name": "Multi",
            "password": "Provisoire-123",
            "all_sites": False,
            "site_ids": [priced.site, priced.site2],
            "roles": [
                {"role_id": viewer_role["id"]},
                {"role_id": export_role["id"], "site_id": priced.site2},
            ],
        },
    )
    assert created.status_code == 201, created.text
    token = login(client, "multi@example.com", "Provisoire-123").json()["access_token"]
    Api(client, token).post(
        "/me/password", json={"current_password": "Provisoire-123", "new_password": PASSWORD}
    )
    token = login(client, "multi@example.com").json()["access_token"]
    multi = Api(client, token)
    _sale(priced)
    depot = _sale(priced, site=priced.site2)
    assert multi.get("/sales").json()["total"] == 2
    # Rôle d'export limité au dépôt : il ne vaut que sur ce site (jamais sur la vue
    # consolidée ni sur le site principal).
    assert _export(multi, "csv").status_code == 403
    on_main = Api(client, token, site_id=uuid.UUID(priced.site))
    assert _export(on_main, "csv").status_code == 403
    on_depot = Api(client, token, site_id=uuid.UUID(priced.site2))
    rows = _csv_rows(_export(on_depot, "csv"))[1:]
    assert [r[0] for r in rows] == [depot["number"]]


def test_export_is_audited(priced: World, client: TestClient) -> None:
    customer = sh.credit_customer(priced)
    _sale(priced, customer_id=customer)
    _sale(priced)
    response = _export(
        priced.owner, "pdf", customer_id=customer, status="VALIDATED", article_reference=" "
    )
    assert response.status_code == 200
    _export(priced.owner, "csv")
    logs = _export_logs(priced)
    assert len(logs) == 2
    latest, first = logs
    assert first["entity_type"] == "export"
    assert first["user_name"] == "Owner alpha" and first["occurred_at"]
    # Filtres réellement renseignés seulement (aucune valeur inventée), format, lignes.
    assert first["data"] == {
        "feature": "sales.history",
        "format": "pdf",
        "filters": {"customer_id": customer, "status": "VALIDATED"},
        "row_count": 1,
    }
    assert latest["data"]["format"] == "csv" and latest["data"]["filters"] == {}
    assert latest["data"]["row_count"] == 2
    # Un export refusé n'est pas tracé comme généré.
    seller = sh.member(priced, client, "vendeur@example.com", "seller", all_sites=True)
    assert _export(seller, "csv").status_code == 403
    assert len(_export_logs(priced)) == 2


def test_export_row_limit(priced: World, settings: Settings, monkeypatch: Any) -> None:
    _sale(priced)
    _sale(priced)
    monkeypatch.setattr(settings, "export_max_rows", 1)
    too_large = _export(priced.owner, "xlsx")
    assert too_large.status_code == 422
    assert too_large.json()["code"] == "export_too_large"
    assert _export_logs(priced) == []
    assert _export(priced.owner, "xlsx", mine=False, status="DRAFT").status_code == 200


# --- Limite de crédit -------------------------------------------------------------------------


def test_credit_limit_requires_dedicated_permission(priced: World, client: TestClient) -> None:
    manager = sh.member(priced, client, "gestion@example.com", "manager", all_sites=True)
    # Créer / modifier un client n'accorde pas la gestion de sa limite.
    denied = manager.post(
        "/customers", json={"customer_type": "INDIVIDUAL", "name": "Awa", "credit_limit": "5000"}
    )
    assert denied.status_code == 403 and denied.json()["code"] == "credit_limit_not_allowed"
    customer = manager.post("/customers", json={"customer_type": "INDIVIDUAL", "name": "Awa"})
    assert customer.status_code == 201
    customer_id = customer.json()["id"]
    changed = manager.patch(f"/customers/{customer_id}", json={"credit_limit": "5000"})
    assert changed.status_code == 403 and changed.json()["code"] == "credit_limit_not_allowed"
    # Valeur inchangée renvoyée par le formulaire : la modification du reste est acceptée.
    renamed = manager.patch(
        f"/customers/{customer_id}", json={"name": "Awa T.", "credit_limit": None}
    )
    assert renamed.status_code == 200 and renamed.json()["name"] == "Awa T."

    # Administrateur : limite modifiée, audit avant / après dédié.
    assert (
        priced.owner.patch(f"/customers/{customer_id}", json={"credit_limit": "5000"})
    ).status_code == 200
    assert (
        priced.owner.patch(f"/customers/{customer_id}", json={"credit_limit": "7500.50"})
    ).status_code == 200
    logs = priced.owner.get(
        "/audit-logs", params={"action": "customer.credit_limit_changed"}
    ).json()["items"]
    assert [(i["data"]["before"], i["data"]["after"]) for i in logs] == [
        ("5000.00", "7500.50"),
        (None, "5000.00"),
    ]
    assert logs[0]["user_name"] == "Owner alpha" and logs[0]["entity_id"] == customer_id
    created = priced.owner.post(
        "/customers", json={"customer_type": "INDIVIDUAL", "name": "Bintou", "credit_limit": "100"}
    )
    assert created.status_code == 201
    logs = priced.owner.get(
        "/audit-logs", params={"action": "customer.credit_limit_changed"}
    ).json()["items"]
    assert (logs[0]["data"]["before"], logs[0]["data"]["after"]) == (None, "100.00")
    # Rôle personnalisé disposant de la permission : autorisé, sans nom de rôle.
    limiter = _custom_member(
        priced,
        client,
        "limite@example.com",
        ["customers.customer.view", "customers.customer.update", "customers.credit_limit.manage"],
    )
    assert (
        limiter.patch(f"/customers/{customer_id}", json={"credit_limit": "0"})
    ).status_code == 200


# --- Fiche enrichie : mouvements et chronologie ---------------------------------------------------


def test_sale_stock_movements_need_movement_permission(priced: World, client: TestClient) -> None:
    sale = _sale(priced, units=3)
    other = _sale(priced, article=1)
    params = {"source_type": "sale", "source_id": sale["id"]}
    movements = priced.owner.get("/stock/movements", params=params).json()
    assert movements["total"] == 1
    item = movements["items"][0]
    assert (item["source_id"], item["quantity"], item["movement_type"]) == (
        sale["id"],
        "-3.000",
        "SALE",
    )
    assert other["id"] not in {m["source_id"] for m in movements["items"]}
    seller = sh.member(priced, client, "vendeur@example.com", "seller", all_sites=True)
    denied = seller.get("/stock/movements", params=params)
    assert denied.status_code == 403
    # Hors périmètre de sites : aucun mouvement.
    depot_viewer = sh.member(
        priced, client, "depot@example.com", "viewer", all_sites=False, site_ids=[priced.site2]
    )
    assert depot_viewer.get("/stock/movements", params=params).json()["total"] == 0


def test_sale_history_needs_audit_permission_and_visible_sale(
    priced: World, client: TestClient
) -> None:
    customer = sh.credit_customer(priced)
    sale = _sale(priced, customer_id=customer, payments=[{"amount": "4000", "method": "CASH"}])
    paid = priced.owner.post(
        f"/sales/{sale['id']}/payments", json={"amount": "6000", "method": "CASH"}
    )
    assert paid.status_code == 201, paid.text
    _sale(priced)  # autre vente : absente de la chronologie
    history = priced.owner.get(f"/sales/{sale['id']}/history")
    assert history.status_code == 200, history.text
    actions = [e["action"] for e in history.json()]
    assert actions[0] == "sale.created" and "sale.validated" in actions
    assert actions.count("payment.created") == 2
    assert all(e["data"].get("sale_id") in (None, sale["id"]) for e in history.json())
    assert history.json()[0]["user_name"] == "Owner alpha"
    # Sans audit.log.view (Gestionnaire) : refusé.
    manager = sh.member(priced, client, "gestion@example.com", "manager", all_sites=True)
    assert manager.get(f"/sales/{sale['id']}/history").status_code == 403
    # audit.log.view sans voir la vente (portée « ses propres ventes ») : introuvable.
    auditor = _custom_member(
        priced, client, "audit@example.com", ["audit.log.view", "sales.sale.view"]
    )
    hidden = auditor.get(f"/sales/{sale['id']}/history")
    assert hidden.status_code == 404 and hidden.json()["code"] == "sale_not_found"
    # Autre entreprise : introuvable.
    missing = priced.owner.get(f"/sales/{uuid.uuid4()}/history")
    assert missing.status_code == 404
